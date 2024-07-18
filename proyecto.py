from tkinter import *
import tkinter.ttk as tk
from tkinter import font
import csv
import textwrap
from fpdf import FPDF
from docx import Document
from pygame import mixer

#Variables globales
root = ""
secondframe = ""
checkbuttons = []
nombres = []
oficios = []
registros = []
nombreV = ""
apellidoV = ""
cargoV = ""
empresaV = ""
calleV = ""
numeroExtV = ""
numeroIntV = ""
coloniaV = ""
municipioV = ""
estadoV = ""
cpV = ""
telefonoV = ""
correoV = ""
nacimientoV = ""
edadV = ""

def text_to_pdf(text, filename):
    a4_width_mm = 210
    pt_to_mm = 0.35
    fontsize_pt = 16
    fontsize_mm = fontsize_pt * pt_to_mm
    margin_bottom_mm = 10
    character_width_mm = 7 * pt_to_mm
    width_text = a4_width_mm / character_width_mm

    pdf = FPDF(orientation='P', unit='mm', format='A4')
    pdf.set_auto_page_break(True, margin=margin_bottom_mm)
    pdf.add_page()
    pdf.set_font(family='Courier', size=fontsize_pt)
    splitted = text.split('\n')

    for line in splitted:
        lines = textwrap.wrap(line, width_text)

        if len(lines) == 0:
            pdf.ln()

        for wrap in lines:
            pdf.cell(0, fontsize_mm, wrap, ln=1)

    pdf.output(filename, 'F')

def createWindow():
    global root
    root = Tk()
    root.title("Combinar correspondencia")
    #w, h  root.winfo_screenwidth(), root.winfo_screenheight()
    root.geometry("1100x600")
    root.resizable(False, False)
    bg = PhotoImage(file="background.gif")
    root.configure(background='skyblue')
    imageLabel = Label(root, image=bg)
    imageLabel.place(x=0, y=0)
    
    # Fuentes
    titlefont = font.Font(family="Comic sans", size=10, weight='bold')
    
    # MainFrame
    
    mainFrame =Frame(root,  width=300,  height=  400,  bg='#d6eaee', borderwidth=2, relief="solid")
    mainFrame.grid(row=0,  column=0,  padx=20,  pady=5)

    global secondFrame
    secondFrame  =  Frame(root,  width=700,  height=400,  bg='#d6eaee', borderwidth=2, relief="solid", padx=20, pady=30)
    secondFrame.grid(row=0,  column=1,  padx=20,  pady=5)
    secondFrame.grid_propagate(False)
    
    thirdFrame = Frame(root,  width=900,  height=150,  bg='#d6eaee', borderwidth=2, relief="solid", padx=10, pady=10)
    thirdFrame.grid(row=1, column=0, columnspan=2, padx=20,  pady=10)
    
    # Widgets
    title = Label(mainFrame, text="Combinar correspondencia", font=titlefont, foreground="black", background='#d6eaee')
    title.grid(row=0, column=0,sticky='n', padx=10, pady=10)
    
    btnGenerarTxt = Button(mainFrame, text="Generar archivo txt", command=createTxt, width=20)
    btnGenerarTxt.grid(row=1, column=0,padx=10, pady=10, sticky=W)
    
    btnGenerarPdf = Button(mainFrame, text="Generar archivo Pdf", command=createPdf, width=20)
    btnGenerarPdf.grid(row=2, column=0,padx=10, pady=10, sticky=W)
    
    btnGenerarWork = Button(mainFrame, text="Generar archivo Word", command=createDocx, width=20)
    btnGenerarWork.grid(row=3, column=0,padx=10, pady=10, sticky=W)
    
    btnEliminar = Button(mainFrame, text="Eliminar registro", command=delete, width=20)
    btnEliminar.grid(row =4, column=0,padx=10, pady=10, sticky=W)
    
    
    #Crear lista
    readCsv()
    for index, nombre in enumerate(nombres):
        registros.append(createRegister(secondFrame,nombre))
        registros[index].grid(column=0, row=index, sticky=(W,E))
    
    Label(thirdFrame, font=titlefont, text="Nombre",foreground="black", background='#d6eaee').grid(column=0, row=0)
    global nombreV 
    nombreV = StringVar()
    nombreEntry = Entry(thirdFrame, width=25, textvariable=nombreV)
    nombreEntry.grid(column=1, row=0)
    Label(thirdFrame, font=titlefont, text="Apellidos",foreground="black", background='#d6eaee').grid(column=2, row=0)
    global apellidoV
    apellidoV = StringVar()
    apellidoEntry = Entry(thirdFrame, width=20, textvariable=apellidoV)
    apellidoEntry.grid(column=3, row=0)
    Label(thirdFrame, font=titlefont, text="Cargo",foreground="black", background='#d6eaee').grid(column=4, row=0)
    global cargoV
    cargoV = StringVar()
    cargoEntry = Entry(thirdFrame, width=20, textvariable=cargoV)
    cargoEntry.grid(column=5, row=0)
    Label(thirdFrame, font=titlefont, text="Empresa",foreground="black", background='#d6eaee').grid(column=6, row=0)
    global empresaV 
    empresaV = StringVar()
    empresaEntry = Entry(thirdFrame, width=25, textvariable=empresaV)
    empresaEntry.grid(column=7, row=0)
    Label(thirdFrame, font=titlefont, text="Calle",foreground="black", background='#d6eaee').grid(column=0, row=1)
    global calleV
    calleV = StringVar()
    calleEntry = Entry(thirdFrame, width=25, textvariable=calleV)
    calleEntry.grid(column=1, row=1)
    Label(thirdFrame, font=titlefont, text="Numero Ext.",foreground="black", background='#d6eaee').grid(column=2, row=1)
    global numeroExtV 
    numeroExtV = StringVar()
    numeroExtEntry = Entry(thirdFrame, width=20, textvariable=numeroExtV)
    numeroExtEntry.grid(column=3, row=1)
    Label(thirdFrame, font=titlefont, text="Numero Int.",foreground="black", background='#d6eaee').grid(column=4, row=1)
    global numeroIntV 
    numeroIntV = StringVar()
    numeroIntEntry = Entry(thirdFrame, width=20, textvariable=numeroIntV)
    numeroIntEntry.grid(column=5, row=1)
    Label(thirdFrame, font=titlefont, text="Colonia",foreground="black", background='#d6eaee').grid(column=6, row=1)
    global coloniaV 
    coloniaV = StringVar()
    coloniaEntry = Entry(thirdFrame, width=25, textvariable=coloniaV)
    coloniaEntry.grid(column=7, row=1)
    Label(thirdFrame, font=titlefont, text="Municipio",foreground="black", background='#d6eaee').grid(column=0, row=2)
    global municipioV 
    municipioV = StringVar()
    municipioEntry = Entry(thirdFrame, width=25, textvariable=municipioV)
    municipioEntry.grid(column=1, row= 2)
    Label(thirdFrame, font=titlefont, text="Estado",foreground="black", background='#d6eaee').grid(column=2, row=2)
    global estadoV 
    estadoV = StringVar()
    estadoEntry = Entry(thirdFrame, width=20, textvariable=estadoV)
    estadoEntry.grid(column=3, row=2)
    Label(thirdFrame, font=titlefont, text="Codigo Postal",foreground="black", background='#d6eaee').grid(column=4, row=2)
    global cpV 
    cpV = StringVar()
    cpEntry = Entry(thirdFrame, width=20, textvariable=cpV)
    cpEntry.grid(column=5, row=2)
    Label(thirdFrame, font=titlefont, text="Telefono",foreground="black", background='#d6eaee').grid(column=6, row=2)
    global telefonoV 
    telefonoV = StringVar()
    telefonoEntry = Entry(thirdFrame, width=25, textvariable=telefonoV)
    telefonoEntry.grid(column=7, row=2)
    Label(thirdFrame, font=titlefont, text="Correo electronico",foreground="black", background='#d6eaee').grid(column=0, row=3)
    global correoV
    correoV = StringVar()
    correoEntry = Entry(thirdFrame, width=25, textvariable=correoV)
    correoEntry.grid(column=1, row=3)
    Label(thirdFrame, font=titlefont, text="Fecha de nacimiento",foreground="black", background='#d6eaee').grid(column=2, row=3)
    global nacimientoV 
    nacimientoV = StringVar()
    nacimientoEntry = Entry(thirdFrame, width=20, textvariable=nacimientoV)
    nacimientoEntry.grid(column=3, row=3)
    Label(thirdFrame, font=titlefont, text="Edad",foreground="black", background='#d6eaee').grid(column=4, row=3)
    global edadV 
    edadV = StringVar()
    edadEntry = Entry(thirdFrame, width=20, textvariable=edadV)
    edadEntry.grid(column=5, row=3)
    
    Button(thirdFrame, text="Agregar", command=lambda: add(secondFrame), width=20, background='#d7d3ef').grid(column=7, row=3)
    
    root.protocol("WM_DELETE_WINDOW", root.quit)
    root.mainloop()
    
def createRegister(parent, labelText):
    global checkbuttons
    newFrame = Frame(parent, padx=0, pady=0,borderwidth=1, relief="solid")
    check = BooleanVar()
    checkbuttons.append(check)
    Checkbutton(newFrame, text="",variable=check,onvalue=1, offvalue=0).grid(column=0, row=0)
    Label(newFrame, text=labelText, width=68).grid(column=1, row=0)
    Button(newFrame, text="Eliminar", width=10, bg='#d7d3ef', command=lambda: delete(labelText)).grid(column=2, row=0)
    Button(newFrame, text="Actualizar", width=10, bg='#d7d3ef', command=lambda: update(labelText)).grid(column=3, row=0)
    return newFrame

def update(nombre):
     update = Toplevel(root, padx=10, pady=10)
     update.title("Actualizar "+nombre)
     Label(update, text="Campo a cambiar",foreground="black", background='#d6eaee').grid(column=0, row=0)
     campoV = StringVar()
     opciones = ["Nombre", "apellidos","cargo","empresa","calle","numeroExt","numeroInt","colonia",
                 "municipio","estado","codigoPostal","telefono","correo","FECHA_NACIMIENTO","edad"]
     campo = tk.Combobox(update, values = opciones, textvariable=campoV)
     campo.grid(column=1, row=0)
     Label(update, text="Nuevo valor",foreground="black", background='#d6eaee').grid(column=0, row=1)
     nuevoV = StringVar()
     nuevo = Entry(update, width=25, textvariable=nuevoV)
     nuevo.grid(column=1, row=1)
     Button(update, text="Actualizar", command=lambda: update2(nombre,campo, nuevoV)).grid(column=0, row=2)     
    # string = f'{nombreV.get()}, {apellidoList[0]}, {apellidoList[1]}, {cargoV.get()}, {empresaV.get()}, {calleV.get()}, {numeroExtV.get()}, {numeroIntV.get()}, {coloniaV.get()}, {municipioV.get()}, {estadoV.get()}, {cpV.get()}, {telefonoV.get()}, {correoV.get()}, {nacimientoV.get()}, {edadV.get()}'+'\n'

def update2(nombre, campo, nuevo):
    str = oficios[nombres.index(nombre)]
    row = str.splitlines()
    #print(str)
    elements = []
    strings = []
    for x, r in enumerate(row):
        if(x<2):continue
        elements = r.split(":")
        strings.append(elements[1])
    i = getTheIndex(campo.get())
    #print(i)
    #delete(nombre)
    strings[i] = nuevo.get()

    string = f'{strings[0]}, {strings[1].split()[0]}, {strings[1].split()[1]}, {strings[2]}, {strings[3]}, {strings[4]}, {strings[5]}, {strings[6]}, {strings[7]}, {strings[8]}, {strings[9]}, {strings[10]}, {strings[11]}, {strings[12]}, {strings[13]}, {strings[14]}'+'\n'
    delete(nombre)
    add2(string)

def getTheIndex(nombre):
    opciones = ["Nombre", "apellidos","cargo","empresa","calle","numeroExt","numeroInt","colonia",
            "municipio","estado","codigoPostal","telefono","correo","FECHA_NACIMIENTO","edad"]
    
    for x, element in enumerate(opciones):
        if(nombre.strip() == element.strip()):return x
    
def delete(nombre):
    index = nombres.index(nombre.strip())
    row = []
    with open("empleados.csv", mode='r') as file:
        row = list(csv.reader(file))
    with open("empleados.csv", mode='w', newline='') as file:
        writer = csv.writer(file, quoting=csv.QUOTE_NONE, escapechar='\\')
        for x,data in enumerate(row):
            if(x==index+1):continue
            cleaned_row_data = [item.rstrip() for item in data]
            writer.writerow(cleaned_row_data)  
    frame = registros[index]
    registros.remove(frame)
    frame.destroy()
    readCsv()

def add2(Cadena):
    global secondFrame

    string =  Cadena
    with open("empleados.csv", mode='a', newline='') as file:
         file.write(string)
    readCsv()
    index = len(registros) 
    registro = createRegister(secondFrame, nombres[index])
    registro.grid(column=0, row=index, sticky=(W, E))
    registros.append(registro)
    
def add(secondFrame):
    apellidoList = apellidoV.get().split()
    if (len(apellidoList)==0):apellidoList.append(" ")
    if (len(apellidoList)==1):apellidoList.append(" ")
    string = f'{nombreV.get()}, {apellidoList[0]}, {apellidoList[1]}, {cargoV.get()}, {empresaV.get()}, {calleV.get()}, {numeroExtV.get()}, {numeroIntV.get()}, {coloniaV.get()}, {municipioV.get()}, {estadoV.get()}, {cpV.get()}, {telefonoV.get()}, {correoV.get()}, {nacimientoV.get()}, {edadV.get()}'+'\n'
    with open("empleados.csv", mode='a', newline='') as file:
         file.write(string)
    readCsv()
    index = len(registros) 
    registro = createRegister(secondFrame, nombres[index])
    registro.grid(column=0, row=index, sticky=(W, E))
    registros.append(registro)

        
def startMusic():
    mixer.init()
    mixer.music.load("Song.mp3")
    mixer.music.play(-1)
    
def readCsv():
    global nombres
    nombres = []
    global oficios
    oficios = []
    with open('empleados.csv', 'r',encoding='utf-8') as file:
        personas = list(csv.DictReader(file))

    with open('oficio.txt', 'r', encoding='utf-8') as oficio_template:
        oficio_template_text = oficio_template.read()
        
    for persona in personas:
        oficio = oficio_template_text
        for key, value in persona.items():
            oficio = oficio.replace(f'[{key.upper()}]', value)
        oficios.append(oficio)
        nombres.append((f"{persona['nombre']}_{persona['apellido1']}").strip(" \t\r"))
    
def getIndex():
    noms= []
    for x,registro in enumerate(registros):
        for widget in registro.winfo_children():
             if(isinstance(widget, Checkbutton)):
                 if checkbuttons[x].get():
                    noms.append(nombres[x])
    return noms

def createTxt():
    index = getIndex()
    for x, k in zip(nombres, oficios):
        if not x in index: continue
        with open(f'txt/{x}_oficio.txt', 'w') as output_file:
            output_file.write(k)

    
def createPdf():
    index = getIndex()
    for x, k in zip(nombres, oficios):
        if not x in index: continue
        final = "pdf/"+x+".pdf"
        text_to_pdf(k,final)
    
def createDocx():
    index = getIndex()
    for x, k in zip(nombres, oficios):
        if not x in index: continue
        document = Document()
        document.add_paragraph(k)
        document.save(f'docx/{x}.docx')
        
if __name__ == "__main__":
    startMusic()
    createWindow()
